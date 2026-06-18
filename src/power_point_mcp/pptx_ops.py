"""Pure-Python helpers around python-pptx.

This module knows nothing about MCP. It exposes plain functions that take
a Presentation (or paths) and return dicts. The server layer wires these
into MCP tools.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.util import Emu

from .config import ServerConfig


def open_target(cfg: ServerConfig) -> PresentationType:
    """Open the bound target presentation."""
    return Presentation(str(cfg.target_path))


def _shape_type_name(shape: Any) -> str:
    st = getattr(shape, "shape_type", None)
    if st is None:
        return "UNKNOWN"
    name = getattr(st, "name", None)
    if name:
        return str(name)
    return str(st)


def _placeholder_type_name(shape: Any) -> str | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    pf = shape.placeholder_format
    pt = getattr(pf, "type", None)
    if pt is None:
        return None
    name = getattr(pt, "name", None)
    return str(name) if name else str(pt)


def _shape_full_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text or ""


def _slide_title(slide: Any) -> str:
    # Best-effort: title placeholder text, falling back to first text frame.
    try:
        title_ph = slide.shapes.title
    except Exception:
        title_ph = None
    if title_ph is not None and getattr(title_ph, "has_text_frame", False):
        text = title_ph.text_frame.text or ""
        if text.strip():
            return text
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            paragraphs = shape.text_frame.paragraphs
            if paragraphs:
                first = paragraphs[0].text or ""
                if first.strip():
                    return first
    return ""


def _slide_text_snippet(slide: Any, limit: int = 120) -> str:
    pieces: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text
            if text:
                pieces.append(text)
    joined = " ".join(p.strip() for p in pieces if p.strip())
    return joined[:limit]


def summarize_presentation(prs: PresentationType) -> dict:
    """Return high-level facts about the presentation."""
    width_emu = int(prs.slide_width) if prs.slide_width is not None else 0
    height_emu = int(prs.slide_height) if prs.slide_height is not None else 0
    layouts = [layout.name for layout in prs.slide_layouts]
    master_name = ""
    if len(prs.slide_masters) > 0:
        master_name = prs.slide_masters[0].name or ""
    has_title = False
    if len(prs.slides) > 0:
        first = prs.slides[0]
        try:
            has_title = first.shapes.title is not None
        except Exception:
            has_title = False
    return {
        "slide_count": len(prs.slides),
        "slide_width_emu": width_emu,
        "slide_height_emu": height_emu,
        "slide_width_inches": float(Emu(width_emu).inches) if width_emu else 0.0,
        "slide_height_inches": float(Emu(height_emu).inches) if height_emu else 0.0,
        "layouts": layouts,
        "master_name": master_name,
        "has_title": has_title,
    }


def list_layouts(prs: PresentationType) -> list[dict]:
    """Return ``[{name, idx, placeholder_names}]`` for every layout in the master."""
    out: list[dict] = []
    for idx, layout in enumerate(prs.slide_layouts):
        placeholder_names = [ph.name or "" for ph in layout.placeholders]
        out.append(
            {
                "name": layout.name or "",
                "idx": idx,
                "placeholder_names": placeholder_names,
            }
        )
    return out


def list_slides(prs: PresentationType) -> list[dict]:
    """Return a compact summary of every slide."""
    out: list[dict] = []
    for index, slide in enumerate(prs.slides):
        layout_name = slide.slide_layout.name if slide.slide_layout else ""
        out.append(
            {
                "index": index,
                "layout_name": layout_name,
                "title": _slide_title(slide),
                "text_snippet": _slide_text_snippet(slide),
            }
        )
    return out


def read_slide(prs: PresentationType, index: int) -> dict:
    """Return shape-level details for a single slide."""
    if index < 0 or index >= len(prs.slides):
        raise IndexError(
            f"slide_index {index} out of range (have {len(prs.slides)} slides)"
        )
    slide = prs.slides[index]
    shapes_out: list[dict] = []
    for shape in slide.shapes:
        is_ph = bool(getattr(shape, "is_placeholder", False))
        ph_idx: int | None = None
        if is_ph:
            try:
                ph_idx = int(shape.placeholder_format.idx)
            except Exception:
                ph_idx = None
        shapes_out.append(
            {
                "shape_id": int(getattr(shape, "shape_id", 0) or 0),
                "name": shape.name or "",
                "shape_type": _shape_type_name(shape),
                "is_placeholder": is_ph,
                "placeholder_idx": ph_idx,
                "placeholder_type": _placeholder_type_name(shape),
                "has_text_frame": bool(getattr(shape, "has_text_frame", False)),
                "text": _shape_full_text(shape),
            }
        )
    return {
        "slide_index": index,
        "layout_name": slide.slide_layout.name if slide.slide_layout else "",
        "shapes": shapes_out,
    }


def create_from_template(
    template_path: Path, target_path: Path, overwrite: bool
) -> None:
    """Open the template and save it as the target file."""
    if target_path.exists() and not overwrite:
        raise FileExistsError(
            f"Target already exists and overwrite=False: {target_path}"
        )
    prs = Presentation(str(template_path))
    target_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target_path))


def add_slide(
    prs: PresentationType,
    layout_name: str,
    placeholders: dict[str, str],
) -> dict:
    """Add a slide using the named layout, filling only supplied placeholders."""
    warnings: list[str] = []

    matches = [layout for layout in prs.slide_layouts if layout.name == layout_name]
    if not matches:
        available = [layout.name for layout in prs.slide_layouts]
        raise KeyError(
            f"No slide layout named {layout_name!r}. Available: {available}"
        )
    if len(matches) > 1:
        warnings.append(
            f"Multiple layouts named {layout_name!r}; using the first."
        )
    layout = matches[0]

    slide = prs.slides.add_slide(layout)
    placeholder_names = {ph.name: ph for ph in slide.placeholders}

    populated: list[str] = []
    for key, value in (placeholders or {}).items():
        ph = placeholder_names.get(key)
        if ph is None:
            warnings.append(f"Unknown placeholder name: {key!r}")
            continue
        if not ph.has_text_frame:
            warnings.append(f"Placeholder {key!r} has no text frame; skipped.")
            continue
        ph.text_frame.text = value
        populated.append(key)

    return {
        "slide_index": len(prs.slides) - 1,
        "layout_name": layout.name,
        "populated": populated,
        "warnings": warnings,
    }


def _sld_id_elements(prs):
    """Return the list of <p:sldId> elements (private API chokepoint)."""
    return list(prs.slides._sldIdLst)


def set_slide_placeholder(
    prs: PresentationType, slide_index: int, placeholder_name: str, text: str
) -> dict:
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise IndexError(f"slide_index {slide_index} out of range")
    slide = prs.slides[slide_index]
    matches = [ph for ph in slide.placeholders if ph.name == placeholder_name]
    if not matches:
        raise KeyError(
            f"no placeholder named {placeholder_name!r} on slide {slide_index}"
        )
    ph = matches[0]
    if not ph.has_text_frame:
        raise ValueError(f"placeholder {placeholder_name!r} has no text frame")
    previous = ph.text_frame.text
    ph.text_frame.text = text
    warnings: list[str] = []
    if len(matches) > 1:
        warnings.append(
            f"multiple placeholders named {placeholder_name!r}; first one used"
        )
    return {
        "slide_index": slide_index,
        "placeholder_name": placeholder_name,
        "placeholder_idx": ph.placeholder_format.idx,
        "previous_text": previous,
        "new_text": text,
        "warnings": warnings,
    }


def set_slide_placeholder_by_idx(
    prs: PresentationType, slide_index: int, placeholder_idx: int, text: str
) -> dict:
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise IndexError(f"slide_index {slide_index} out of range")
    slide = prs.slides[slide_index]
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == placeholder_idx:
            if not ph.has_text_frame:
                raise ValueError(
                    f"placeholder idx {placeholder_idx} has no text frame"
                )
            previous = ph.text_frame.text
            ph.text_frame.text = text
            return {
                "slide_index": slide_index,
                "placeholder_idx": placeholder_idx,
                "placeholder_name": ph.name,
                "previous_text": previous,
                "new_text": text,
            }
    raise KeyError(
        f"no placeholder with idx {placeholder_idx} on slide {slide_index}"
    )


def set_slide_title(prs: PresentationType, slide_index: int, text: str) -> dict:
    if slide_index < 0 or slide_index >= len(prs.slides):
        raise IndexError(f"slide_index {slide_index} out of range")
    slide = prs.slides[slide_index]
    title = slide.shapes.title
    if title is None:
        # fallback: placeholder with idx 0
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 0:
                title = ph
                break
    if title is None or not title.has_text_frame:
        raise KeyError(f"slide {slide_index} has no title placeholder")
    previous = title.text_frame.text
    title.text_frame.text = text
    return {
        "slide_index": slide_index,
        "placeholder_name": title.name,
        "previous_text": previous,
        "new_text": text,
    }


def delete_slide(prs: PresentationType, slide_index: int) -> dict:
    elements = _sld_id_elements(prs)
    if slide_index < 0 or slide_index >= len(elements):
        raise IndexError(f"slide_index {slide_index} out of range")
    sld_id_lst = prs.slides._sldIdLst
    target = elements[slide_index]
    # drop the relationship from the presentation part
    rId = target.attrib.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    sld_id_lst.remove(target)
    if rId:
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
    return {"deleted_index": slide_index, "remaining_slide_count": len(prs.slides)}


def reorder_slide(
    prs: PresentationType, slide_index: int, new_index: int
) -> dict:
    elements = _sld_id_elements(prs)
    n = len(elements)
    if slide_index < 0 or slide_index >= n:
        raise IndexError(f"slide_index {slide_index} out of range")
    # clamp new_index
    new_index = max(0, min(new_index, n - 1))
    if new_index == slide_index:
        return {"old_index": slide_index, "new_index": new_index, "slide_count": n}
    sld_id_lst = prs.slides._sldIdLst
    target = elements[slide_index]
    sld_id_lst.remove(target)
    # re-grab elements after removal to find the right insertion point
    remaining = list(sld_id_lst)
    if new_index >= len(remaining):
        sld_id_lst.append(target)
    else:
        sld_id_lst.insert(list(sld_id_lst).index(remaining[new_index]), target)
    return {"old_index": slide_index, "new_index": new_index, "slide_count": n}
