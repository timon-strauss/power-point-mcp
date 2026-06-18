"""MCP server wiring. Each tool re-opens the bound presentation for fresh state."""

from __future__ import annotations

import functools
import inspect
from typing import Any

from mcp.server.fastmcp import FastMCP
from pptx import Presentation

from . import pptx_ops
from .config import ConfigError, ServerConfig
from .security import SecurityError, assert_within_target


def _safe(fn):
    """Wrap a tool body so structured errors are returned, not raised.

    Preserves the wrapped function's signature so FastMCP/pydantic can
    introspect parameters correctly (otherwise the wrapper's ``*args,
    **kwargs`` would surface as required fields).
    """

    @functools.wraps(fn)
    def wrapper(*args: Any, **kwargs: Any) -> Any:
        try:
            return fn(*args, **kwargs)
        except (
            ConfigError,
            SecurityError,
            FileNotFoundError,
            FileExistsError,
            IndexError,
            KeyError,
            ValueError,
        ) as exc:
            return {"error": f"{type(exc).__name__}: {exc}"}

    wrapper.__signature__ = inspect.signature(fn)  # type: ignore[attr-defined]
    wrapper.__annotations__ = dict(getattr(fn, "__annotations__", {}))
    return wrapper


def create_server(cfg: ServerConfig) -> FastMCP:
    """Build the FastMCP server bound to ``cfg``."""
    mcp = FastMCP("power-point-mcp")

    def _open() -> Any:
        path = assert_within_target(cfg.target_path, cfg)
        if not path.exists():
            raise FileNotFoundError(
                f"Bound target does not exist yet: {path}. "
                "Call create_presentation_from_template first."
            )
        return Presentation(str(path))

    @mcp.tool()
    @_safe
    def presentation_info() -> dict:
        """Return high-level facts about the bound presentation."""
        prs = _open()
        return pptx_ops.summarize_presentation(prs)

    @mcp.tool(name="list_slides")
    @_safe
    def list_slides_tool() -> list[dict]:
        """List every slide with its layout, title, and a short text snippet."""
        prs = _open()
        return pptx_ops.list_slides(prs)

    @mcp.tool(name="list_layouts")
    @_safe
    def list_layouts_tool() -> list[dict]:
        """List every layout in the master with its idx and placeholder names."""
        prs = _open()
        return pptx_ops.list_layouts(prs)

    @mcp.tool(name="read_slide")
    @_safe
    def read_slide_tool(slide_index: int) -> dict:
        """Return shape-level details for a single slide."""
        prs = _open()
        return pptx_ops.read_slide(prs, slide_index)

    @mcp.tool()
    @_safe
    def create_presentation_from_template(overwrite: bool = False) -> dict:
        """Create the bound target file from PPTX_TEMPLATE."""
        if cfg.template_path is None:
            raise ConfigError(
                "PPTX_TEMPLATE is not set; cannot create from template."
            )
        target = assert_within_target(cfg.target_path, cfg)
        pptx_ops.create_from_template(cfg.template_path, target, overwrite)
        return {
            "created": str(target),
            "template": str(cfg.template_path),
            "overwrite": overwrite,
        }

    @mcp.tool(name="add_slide")
    @_safe
    def add_slide_tool(
        layout_name: str,
        placeholders: dict[str, str] | None = None,
    ) -> dict:
        """Add a slide using ``layout_name``; fill ONLY supplied placeholders."""
        target = assert_within_target(cfg.target_path, cfg)
        if not target.exists():
            raise FileNotFoundError(
                f"Bound target does not exist yet: {target}. "
                "Call create_presentation_from_template first."
            )
        prs = Presentation(str(target))
        result = pptx_ops.add_slide(prs, layout_name, placeholders or {})
        prs.save(str(target))
        return result

    def _open_for_write() -> tuple[Any, Any]:
        target = assert_within_target(cfg.target_path, cfg)
        if not target.exists():
            raise FileNotFoundError(
                f"Bound target does not exist yet: {target}. "
                "Call create_presentation_from_template first."
            )
        return target, Presentation(str(target))

    @mcp.tool(name="set_slide_placeholder")
    @_safe
    def set_slide_placeholder_tool(
        slide_index: int, placeholder_name: str, text: str
    ) -> dict:
        """Set a placeholder on an existing slide by its name."""
        target, prs = _open_for_write()
        result = pptx_ops.set_slide_placeholder(
            prs, slide_index, placeholder_name, text
        )
        prs.save(str(target))
        return result

    @mcp.tool(name="set_slide_placeholder_by_idx")
    @_safe
    def set_slide_placeholder_by_idx_tool(
        slide_index: int, placeholder_idx: int, text: str
    ) -> dict:
        """Set a placeholder on an existing slide by its layout idx."""
        target, prs = _open_for_write()
        result = pptx_ops.set_slide_placeholder_by_idx(
            prs, slide_index, placeholder_idx, text
        )
        prs.save(str(target))
        return result

    @mcp.tool(name="set_slide_title")
    @_safe
    def set_slide_title_tool(slide_index: int, text: str) -> dict:
        """Set the title placeholder of a slide."""
        target, prs = _open_for_write()
        result = pptx_ops.set_slide_title(prs, slide_index, text)
        prs.save(str(target))
        return result

    @mcp.tool(name="delete_slide")
    @_safe
    def delete_slide_tool(slide_index: int) -> dict:
        """Remove a slide from the bound presentation."""
        target, prs = _open_for_write()
        result = pptx_ops.delete_slide(prs, slide_index)
        prs.save(str(target))
        return result

    @mcp.tool(name="reorder_slide")
    @_safe
    def reorder_slide_tool(slide_index: int, new_index: int) -> dict:
        """Move a slide to a new position."""
        target, prs = _open_for_write()
        result = pptx_ops.reorder_slide(prs, slide_index, new_index)
        prs.save(str(target))
        return result

    return mcp
