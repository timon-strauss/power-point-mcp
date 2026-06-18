"""Tests for delete_slide and reorder_slide structural mutations."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from power_point_mcp import pptx_ops


def _build_three_titled_slides(target: Path, tmp_pptx_template: Path) -> Path:
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)
    prs = Presentation(str(target))

    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    title_ph_name = next(
        ph.name for ph in layout.placeholders if ph.placeholder_format.idx == 0
    )

    for label in ("A", "B", "C"):
        pptx_ops.add_slide(prs, "Title Slide", {title_ph_name: label})
    prs.save(str(target))
    return target


def _slide_titles(prs) -> list[str]:
    titles: list[str] = []
    for index in range(len(prs.slides)):
        detail = pptx_ops.read_slide(prs, index)
        for shape in detail["shapes"]:
            if (
                shape["is_placeholder"]
                and shape["placeholder_idx"] == 0
                and shape["text"]
            ):
                titles.append(shape["text"])
                break
        else:
            titles.append("")
    return titles


def test_delete_slide_middle(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    result = pptx_ops.delete_slide(prs, 1)
    assert result == {"deleted_index": 1, "remaining_slide_count": 2}
    prs.save(str(target))

    prs2 = Presentation(str(target))
    titles = _slide_titles(prs2)
    assert titles == ["A", "C"]


def test_delete_slide_first(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    pptx_ops.delete_slide(prs, 0)
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["B", "C"]


def test_delete_slide_last(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    pptx_ops.delete_slide(prs, 2)
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["A", "B"]


def test_delete_slide_out_of_range(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    with pytest.raises(IndexError):
        pptx_ops.delete_slide(prs, 99)
    with pytest.raises(IndexError):
        pptx_ops.delete_slide(prs, -1)


def test_reorder_slide_to_front(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    result = pptx_ops.reorder_slide(prs, 2, 0)
    assert result == {"old_index": 2, "new_index": 0, "slide_count": 3}
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["C", "A", "B"]


def test_reorder_slide_to_end(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    pptx_ops.reorder_slide(prs, 0, 2)
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["B", "C", "A"]


def test_reorder_slide_clamps_new_index(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    # new_index = 99 should clamp to last
    result = pptx_ops.reorder_slide(prs, 0, 99)
    assert result["new_index"] == 2
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["B", "C", "A"]


def test_reorder_slide_noop(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    result = pptx_ops.reorder_slide(prs, 1, 1)
    assert result == {"old_index": 1, "new_index": 1, "slide_count": 3}
    prs.save(str(target))

    prs2 = Presentation(str(target))
    assert _slide_titles(prs2) == ["A", "B", "C"]


def test_reorder_slide_out_of_range(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = _build_three_titled_slides(tmp_path / "out.pptx", tmp_pptx_template)
    prs = Presentation(str(target))
    with pytest.raises(IndexError):
        pptx_ops.reorder_slide(prs, 99, 0)
    with pytest.raises(IndexError):
        pptx_ops.reorder_slide(prs, -1, 0)
