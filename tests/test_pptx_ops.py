"""Tests for power_point_mcp.pptx_ops."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from power_point_mcp import pptx_ops


def _layouts(path: Path) -> list[str]:
    return [layout.name for layout in Presentation(str(path)).slide_layouts]


def test_create_from_template_copies_layouts(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=False)
    assert target.exists()
    assert _layouts(target) == _layouts(tmp_pptx_template)


def test_create_from_template_no_overwrite(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=False)
    with pytest.raises(FileExistsError):
        pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=False)
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)


def test_summarize_zero_then_one_slide(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)

    prs = Presentation(str(target))
    info = pptx_ops.summarize_presentation(prs)
    assert info["slide_count"] == 0
    assert "Title Slide" in info["layouts"]

    pptx_ops.add_slide(prs, "Title Slide", {})
    prs.save(str(target))

    prs2 = Presentation(str(target))
    info2 = pptx_ops.summarize_presentation(prs2)
    assert info2["slide_count"] == 1


def test_add_slide_fills_only_supplied_placeholders(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)
    prs = Presentation(str(target))

    # Find the actual placeholder name on the Title Slide layout.
    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    title_ph_name = None
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == 0:
            title_ph_name = ph.name
            break
    assert title_ph_name is not None

    result = pptx_ops.add_slide(prs, "Title Slide", {title_ph_name: "Hello"})
    assert result["populated"] == [title_ph_name]
    assert result["warnings"] == []

    result2 = pptx_ops.add_slide(
        prs, "Title Slide", {title_ph_name: "Hi", "Bogus": "X"}
    )
    assert title_ph_name in result2["populated"]
    assert any("Bogus" in w for w in result2["warnings"])


def test_read_slide_returns_title_text(
    tmp_path: Path, tmp_pptx_template: Path
) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)
    prs = Presentation(str(target))

    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    title_ph_name = next(
        ph.name for ph in layout.placeholders if ph.placeholder_format.idx == 0
    )
    pptx_ops.add_slide(prs, "Title Slide", {title_ph_name: "Hello there"})
    prs.save(str(target))

    prs2 = Presentation(str(target))
    detail = pptx_ops.read_slide(prs2, 0)
    title_texts = [s["text"] for s in detail["shapes"] if s["is_placeholder"]]
    assert any("Hello there" in t for t in title_texts)


def test_read_slide_out_of_range(tmp_path: Path, tmp_pptx_template: Path) -> None:
    target = tmp_path / "out.pptx"
    pptx_ops.create_from_template(tmp_pptx_template, target, overwrite=True)
    prs = Presentation(str(target))
    with pytest.raises(IndexError):
        pptx_ops.read_slide(prs, 0)
    with pytest.raises(IndexError):
        pptx_ops.read_slide(prs, 99)
