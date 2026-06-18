"""End-to-end-ish tests that drive the registered tool functions directly.

We bypass the transport layer because exercising stdio inside pytest is
fragile. The tools are still real FastMCP-registered callables; this test
verifies the wiring (open -> mutate -> save -> reopen) works.
"""

from __future__ import annotations

import asyncio
from typing import Any

import pytest
from pptx import Presentation

from power_point_mcp.config import ServerConfig
from power_point_mcp.server import create_server


def _call(server: Any, name: str, **kwargs: Any) -> Any:
    """Invoke a registered FastMCP tool by name and return its python value."""
    result = asyncio.run(server.call_tool(name, kwargs))
    # FastMCP.call_tool returns either (content_list, structured_value) on
    # newer versions or just content_list. Prefer the structured value.
    if isinstance(result, tuple) and len(result) == 2:
        result = result[1]
    # FastMCP wraps non-object tool returns (and sometimes object returns)
    # in a ``{"result": ...}`` envelope for the structured payload. Unwrap
    # it so callers see the tool's actual return value.
    if isinstance(result, dict) and set(result.keys()) == {"result"}:
        return result["result"]
    return result


def test_full_flow(cfg: ServerConfig) -> None:
    server = create_server(cfg)

    created = _call(server, "create_presentation_from_template", overwrite=True)
    assert isinstance(created, dict)
    assert "error" not in created

    info = _call(server, "presentation_info")
    assert isinstance(info, dict)
    assert info.get("slide_count") == 0

    # Discover the actual title placeholder name from the freshly created file.
    prs = Presentation(str(cfg.target_path))
    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    title_ph_name = next(
        ph.name for ph in layout.placeholders if ph.placeholder_format.idx == 0
    )

    added = _call(
        server,
        "add_slide",
        layout_name="Title Slide",
        placeholders={title_ph_name: "Hi"},
    )
    assert isinstance(added, dict)
    assert added.get("slide_index") == 0

    listed = _call(server, "list_slides")
    assert isinstance(listed, list)
    assert len(listed) == 1
    assert listed[0]["layout_name"] == "Title Slide"


def test_add_slide_before_create_returns_error(cfg: ServerConfig) -> None:
    server = create_server(cfg)
    out = _call(server, "add_slide", layout_name="Title Slide", placeholders={})
    assert isinstance(out, dict)
    assert "error" in out


def test_create_from_template_requires_template(tmp_path: Any) -> None:
    cfg = ServerConfig(
        target_path=(tmp_path / "out.pptx").resolve(), template_path=None
    )
    server = create_server(cfg)
    out = _call(server, "create_presentation_from_template", overwrite=True)
    assert isinstance(out, dict)
    assert "error" in out


def _setup_with_two_slides(server: Any, cfg: ServerConfig) -> str:
    """Create the file and add two Title-Slide slides; return the title ph name."""
    _call(server, "create_presentation_from_template", overwrite=True)
    prs = Presentation(str(cfg.target_path))
    layout = next(l for l in prs.slide_layouts if l.name == "Title Slide")
    title_ph_name = next(
        ph.name for ph in layout.placeholders if ph.placeholder_format.idx == 0
    )
    _call(
        server,
        "add_slide",
        layout_name="Title Slide",
        placeholders={title_ph_name: "First"},
    )
    _call(
        server,
        "add_slide",
        layout_name="Title Slide",
        placeholders={title_ph_name: "Second"},
    )
    return title_ph_name


def test_set_slide_title_e2e(cfg: ServerConfig) -> None:
    server = create_server(cfg)
    _setup_with_two_slides(server, cfg)

    out = _call(server, "set_slide_title", slide_index=0, text="Renamed")
    assert isinstance(out, dict)
    assert "error" not in out
    assert out["new_text"] == "Renamed"

    listed = _call(server, "list_slides")
    assert listed[0]["title"] == "Renamed"


def test_delete_slide_e2e(cfg: ServerConfig) -> None:
    server = create_server(cfg)
    _setup_with_two_slides(server, cfg)

    out = _call(server, "delete_slide", slide_index=0)
    assert isinstance(out, dict)
    assert out.get("remaining_slide_count") == 1

    listed = _call(server, "list_slides")
    assert len(listed) == 1
    assert listed[0]["title"] == "Second"


def test_reorder_slide_e2e(cfg: ServerConfig) -> None:
    server = create_server(cfg)
    _setup_with_two_slides(server, cfg)

    out = _call(server, "reorder_slide", slide_index=0, new_index=1)
    assert isinstance(out, dict)
    assert out.get("new_index") == 1

    listed = _call(server, "list_slides")
    assert [s["title"] for s in listed] == ["Second", "First"]


def test_set_slide_placeholder_bad_index_returns_error(cfg: ServerConfig) -> None:
    server = create_server(cfg)
    _call(server, "create_presentation_from_template", overwrite=True)
    out = _call(
        server,
        "set_slide_placeholder",
        slide_index=99,
        placeholder_name="x",
        text="y",
    )
    assert isinstance(out, dict)
    assert "error" in out
